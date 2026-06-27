"""
Analytics Export (ARB #42)
--------------------------
Turn a rendered analytics chart (a ``ChartSpec`` — the same payload the Studio /
Chart Builder draw with) into a downloadable **PDF** or **PowerPoint** deck.

  • PDF  — reportlab: title block + a native vector chart + the underlying data
    table. Mirrors the existing ``analytics_service.roi_pdf_bytes`` style.
  • PPTX — python-pptx: one slide with a *native, editable* chart built from the
    series, so stakeholders can restyle it in PowerPoint.

Both take the spec the frontend already holds, so anything you can see you can
export. Input is treated defensively (best-effort coercion) — a malformed series
yields an empty-but-valid file rather than a 500.
"""

from __future__ import annotations

import io
import datetime
from typing import Any

# Brand palette (kept consistent with document_generation/generator.py)
_BRAND_DARK = (0x0A, 0x25, 0x40)
_BRAND_CYAN = (0x00, 0xA2, 0x9A)
_SERIES_HEX = ["#00A29A", "#6366F1", "#F59E0B", "#EF4444", "#10B981", "#8B5CF6", "#0EA5E9"]


def _rows(spec: dict) -> list[dict]:
    data = spec.get("data") or []
    return [r for r in data if isinstance(r, dict)]


def _y_keys(spec: dict) -> list[str]:
    yk = spec.get("y_keys") or []
    return [k for k in yk if k] or ["value"]


def _y_label(spec: dict, key: str) -> str:
    return (spec.get("y_labels") or {}).get(key, key)


def _num(v: Any) -> float:
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


# ── PowerPoint ────────────────────────────────────────────────────────────────

def chart_to_pptx(spec: dict) -> bytes:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION

    type_map = {
        "bar": XL_CHART_TYPE.COLUMN_CLUSTERED,
        "line": XL_CHART_TYPE.LINE_MARKERS,
        "area": XL_CHART_TYPE.AREA,
        "pie": XL_CHART_TYPE.PIE,
    }
    chart_type = type_map.get(spec.get("type"), XL_CHART_TYPE.COLUMN_CLUSTERED)

    rows = _rows(spec)
    x_key = spec.get("x_key") or "label"
    y_keys = _y_keys(spec)
    categories = [str(r.get(x_key, "—")) for r in rows]

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[5])  # title-only

    # Title
    slide.shapes.title.text = spec.get("title") or "Analytics Chart"
    if spec.get("subtitle"):
        box = slide.shapes.add_textbox(Inches(0.5), Inches(1.15), Inches(12.3), Inches(0.5))
        p = box.text_frame.paragraphs[0]
        p.text = str(spec["subtitle"])
        p.font.size = Pt(14)

    chart_data = CategoryChartData()
    chart_data.categories = categories or ["—"]
    if chart_type == XL_CHART_TYPE.PIE:
        # Pie shows a single series only.
        key = y_keys[0]
        chart_data.add_series(_y_label(spec, key), [_num(r.get(key)) for r in rows] or [0])
    else:
        for key in y_keys:
            chart_data.add_series(_y_label(spec, key), [_num(r.get(key)) for r in rows] or [0])

    gf = slide.shapes.add_chart(
        chart_type, Inches(0.6), Inches(1.7), Inches(12.1), Inches(5.4), chart_data
    )
    chart = gf.chart
    chart.has_legend = len(y_keys) > 1 or chart_type == XL_CHART_TYPE.PIE
    if chart.has_legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


# ── PDF ───────────────────────────────────────────────────────────────────────

def chart_to_pdf(spec: dict) -> bytes:
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.units import mm
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import (
        SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle,
    )
    from reportlab.graphics.shapes import Drawing
    from reportlab.graphics.charts.barcharts import VerticalBarChart
    from reportlab.graphics.charts.linecharts import HorizontalLineChart
    from reportlab.graphics.charts.piecharts import Pie
    from reportlab.graphics.charts.legends import Legend

    rows = _rows(spec)
    x_key = spec.get("x_key") or "label"
    y_keys = _y_keys(spec)
    unit = spec.get("unit") or ""
    categories = [str(r.get(x_key, "—")) for r in rows]
    brand = colors.Color(*[c / 255 for c in _BRAND_DARK])
    cyan = colors.Color(*[c / 255 for c in _BRAND_CYAN])
    palette = [colors.HexColor(h) for h in _SERIES_HEX]

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=landscape(A4),
        topMargin=16 * mm, bottomMargin=16 * mm, leftMargin=18 * mm, rightMargin=18 * mm,
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "ChartTitle", parent=styles["Title"], textColor=brand, fontSize=20, spaceAfter=2,
    )
    sub_style = ParagraphStyle(
        "ChartSub", parent=styles["Normal"], textColor=colors.grey, fontSize=10, spaceAfter=10,
    )
    elems: list = [Paragraph(spec.get("title") or "Analytics Chart", title_style)]
    sub = spec.get("subtitle") or ""
    stamp = datetime.datetime.now().strftime("%d %b %Y, %H:%M")
    elems.append(Paragraph(f"{sub + ' · ' if sub else ''}Generated {stamp}", sub_style))

    # ── Native vector chart ──
    drawing = Drawing(700, 280)
    chart_type = spec.get("type")
    try:
        if chart_type == "pie" and rows:
            key = y_keys[0]
            pie = Pie()
            pie.x, pie.y, pie.width, pie.height = 230, 30, 220, 220
            pie.data = [_num(r.get(key)) for r in rows]
            pie.labels = categories
            for i in range(len(pie.data)):
                pie.slices[i].fillColor = palette[i % len(palette)]
            drawing.add(pie)
        elif chart_type == "line" and rows:
            lc = HorizontalLineChart()
            lc.x, lc.y, lc.width, lc.height = 60, 50, 600, 200
            lc.data = [[_num(r.get(k)) for r in rows] for k in y_keys]
            lc.categoryAxis.categoryNames = categories
            lc.categoryAxis.labels.angle = 30
            lc.categoryAxis.labels.dy = -8
            for i in range(len(y_keys)):
                lc.lines[i].strokeColor = palette[i % len(palette)]
                lc.lines[i].strokeWidth = 2
            drawing.add(lc)
        elif rows:
            bc = VerticalBarChart()
            bc.x, bc.y, bc.width, bc.height = 60, 50, 600, 200
            bc.data = [[_num(r.get(k)) for r in rows] for k in y_keys]
            bc.categoryAxis.categoryNames = categories
            bc.categoryAxis.labels.angle = 30
            bc.categoryAxis.labels.dy = -8
            bc.valueAxis.valueMin = 0
            for i in range(len(y_keys)):
                bc.bars[i].fillColor = palette[i % len(palette)]
            drawing.add(bc)
        elif len(y_keys) > 1:
            legend = Legend()
            legend.x, legend.y = 480, 250
            legend.colorNamePairs = [
                (palette[i % len(palette)], _y_label(spec, k)) for i, k in enumerate(y_keys)
            ]
            drawing.add(legend)
    except Exception:
        # Charting is decorative — never let it sink the export. Table still renders.
        drawing = Drawing(700, 10)
    elems.append(drawing)
    elems.append(Spacer(1, 8 * mm))

    # ── Data table ──
    header = [x_key] + [f"{_y_label(spec, k)}{f' ({unit})' if unit else ''}" for k in y_keys]
    table_rows = [header]
    for r in rows:
        table_rows.append([str(r.get(x_key, "—"))] + [f"{_num(r.get(k)):g}" for k in y_keys])
    if len(table_rows) == 1:
        table_rows.append(["No data", *["—"] * len(y_keys)])

    table = Table(table_rows, repeatRows=1, hAlign="LEFT")
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), brand),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.Color(0.96, 0.98, 0.98)]),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.lightgrey),
        ("LINEBELOW", (0, 0), (-1, 0), 1, cyan),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ("ALIGN", (1, 0), (-1, -1), "RIGHT"),
    ]))
    elems.append(table)

    doc.build(elems)
    return buf.getvalue()
