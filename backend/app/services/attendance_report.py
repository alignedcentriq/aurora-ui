"""
Attendance Report Builder
-------------------------
Turns a `attendance_service.team_report(...)` dict into a branded .xlsx workbook (one row
per employee + a totals row) and returns it base64-encoded so it can be attached to a
Microsoft Graph email (see email_service._send `files=`).

openpyxl is already a project dependency; no new packages needed.
"""

import base64
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

# Brand palette (matches the Gradient Hero email header)
_HEADER_FILL = PatternFill("solid", fgColor="1B6FC8")
_TOTAL_FILL = PatternFill("solid", fgColor="E2E8F0")
_HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
_TOTAL_FONT = Font(bold=True, color="0F172A")
_THIN = Side(style="thin", color="CBD5E1")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)

_COLUMNS = [
    ("Employee", "employee", 26),
    ("Email", "email", 30),
    ("Department", "department", 20),
    ("Designation", "designation", 22),
    ("Reports To", "reports_to", 22),
    ("Present", "present", 10),
    ("Absent", "absent", 10),
    ("WFH", "wfh", 10),
    ("Late", "late", 10),
    ("Half-day", "half_day", 11),
]


def build_team_xlsx(report: dict) -> tuple[str, str]:
    """
    Build the team attendance workbook from a successful team_report dict.

    Returns (filename, base64_contents). Raises ValueError if the report isn't a success
    payload (caller should guard on report["success"] first).
    """
    if not report.get("success"):
        raise ValueError("build_team_xlsx requires a successful team_report payload")

    wb = Workbook()
    ws = wb.active
    ws.title = "Attendance"

    # Title band
    last_col = get_column_letter(len(_COLUMNS))
    ws.merge_cells(f"A1:{last_col}1")
    title = ws["A1"]
    title.value = f"Team Attendance — {report['period']}"
    title.font = Font(bold=True, size=14, color="0F172A")
    title.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 24

    ws.merge_cells(f"A2:{last_col}2")
    sub = ws["A2"]
    sub.value = f"Manager: {report['manager']}  •  {report['headcount']} people in hierarchy"
    sub.font = Font(size=10, color="475569")

    header_row = 4
    for c, (label, _key, width) in enumerate(_COLUMNS, start=1):
        cell = ws.cell(row=header_row, column=c, value=label)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center" if c > 5 else "left", vertical="center")
        cell.border = _BORDER
        ws.column_dimensions[get_column_letter(c)].width = width

    r = header_row + 1
    for m in report["members"]:
        for c, (_label, key, _w) in enumerate(_COLUMNS, start=1):
            val = m.get(key, "")
            cell = ws.cell(row=r, column=c, value=val)
            cell.border = _BORDER
            cell.alignment = Alignment(horizontal="center" if c > 5 else "left")
        r += 1

    # Totals row
    totals = report["totals"]
    tcell = ws.cell(row=r, column=1, value="TOTAL")
    tcell.font = _TOTAL_FONT
    tcell.fill = _TOTAL_FILL
    tcell.border = _BORDER
    for c in range(2, 6):
        cell = ws.cell(row=r, column=c, value="")
        cell.fill = _TOTAL_FILL
        cell.border = _BORDER
    for c, (_label, key, _w) in enumerate(_COLUMNS, start=1):
        if key in totals:
            cell = ws.cell(row=r, column=c, value=totals[key])
            cell.font = _TOTAL_FONT
            cell.fill = _TOTAL_FILL
            cell.border = _BORDER
            cell.alignment = Alignment(horizontal="center")

    ws.freeze_panes = f"A{header_row + 1}"

    buf = io.BytesIO()
    wb.save(buf)
    b64 = base64.b64encode(buf.getvalue()).decode("ascii")

    safe_period = report["period"].replace(" ", "_")
    filename = f"Team_Attendance_{safe_period}.xlsx"
    return filename, b64
